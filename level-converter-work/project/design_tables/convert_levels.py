from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print(
        "Missing dependency: openpyxl\n"
        "Run: py -3 -m pip install -r design_tables/requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1)


TOOL_VERSION = 3
MAP_SHEET_PREFIX = "map_"
GRID_START_ROW = 9
GRID_START_COLUMN = 2
BASE_CODES = {"G", "R", "W", "BR", "X"}
STATIC_CODES = {"T", "BU", "GT", "H1", "H2"}
SINGLE_CELL_CODES = {"RB", "P", "C", "N", "E", "BO"}
OBJECT_CODES = STATIC_CODES | SINGLE_CELL_CODES
TRUE_VALUES = {"true", "yes", "y", "1", "是"}
FALSE_VALUES = {"false", "no", "n", "0", "否", ""}
ENEMY_RESOURCE_SIZES = {
    "bat": (16, 24),
    "male_cow_brown": (32, 32),
    "female_cow_brown": (32, 32),
    "chicken_red": (16, 16),
    "chicken_blonde_green": (16, 16),
    "baby_chicken_yellow": (16, 16),
}


class ConversionError(RuntimeError):
    pass


@dataclass(frozen=True)
class LegendEntry:
    code: str
    resource_path: str
    footprint_width: int
    footprint_height: int
    passability: str


@dataclass(frozen=True)
class Placement:
    code: str
    token: str
    x: int
    y: int
    width: int = 1
    height: int = 1


@dataclass
class MapSpec:
    sheet_name: str
    map_id: str
    map_name: str
    width: int
    height: int
    cell_size: int
    default_terrain: str
    output_scene: str
    is_default: bool
    grid: list[list[str]]
    placements: list[Placement]
    spawn: Placement


def parse_args() -> argparse.Namespace:
    default_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(
        description="Convert levels.xlsx map sheets into Godot scenes."
    )
    parser.add_argument("--project-root", type=Path, default=default_root)
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--workbook", type=Path)
    return parser.parse_args()


def fail(sheet: str, cell: str, message: str) -> None:
    raise ConversionError(f"[{sheet}!{cell}] {message}")


def cell_name(x: int, y: int) -> str:
    column = get_column_letter(GRID_START_COLUMN + x)
    return f"{column}{GRID_START_ROW + y}"


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_positive_int(sheet: str, cell: str, value: Any) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        fail(sheet, cell, "must be an integer")
    if parsed <= 0:
        fail(sheet, cell, "must be greater than zero")
    return parsed


def parse_default_flag(sheet: str, value: Any) -> bool:
    normalized = clean_text(value).lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    fail(sheet, "L4", "default flag must be yes/no or true/false")


def parse_legend(workbook: Any, project_root: Path) -> dict[str, LegendEntry]:
    if "图例" not in workbook.sheetnames:
        raise ConversionError("Workbook is missing the legend sheet")
    sheet = workbook["图例"]
    entries: dict[str, LegendEntry] = {}
    blank_rows = 0
    rows = sheet.iter_rows(min_row=5, max_row=255, max_col=18)
    for row, cells in enumerate(rows, start=5):
        code = clean_text(cells[0].value)
        if not code:
            blank_rows += 1
            if blank_rows >= 8:
                break
            continue
        blank_rows = 0
        if code in entries:
            fail(sheet.title, f"A{row}", f"duplicate legend code '{code}'")
        original_width = parse_positive_int(
            sheet.title, f"F{row}", cells[5].value
        )
        original_height = parse_positive_int(
            sheet.title, f"G{row}", cells[6].value
        )
        try:
            display_scale = float(cells[7].value)
        except (TypeError, ValueError):
            fail(sheet.title, f"H{row}", "display scale must be numeric")
        if display_scale <= 0:
            fail(sheet.title, f"H{row}", "display scale must be greater than zero")
        override_width = cells[12].value
        override_height = cells[13].value
        footprint_width = (
            parse_positive_int(sheet.title, f"M{row}", override_width)
            if override_width not in (None, "")
            else max(1, math.ceil(original_width * display_scale / 32.0))
        )
        footprint_height = (
            parse_positive_int(sheet.title, f"N{row}", override_height)
            if override_height not in (None, "")
            else max(1, math.ceil(original_height * display_scale / 32.0))
        )
        resource_path = clean_text(cells[4].value)
        if resource_path.startswith("res://"):
            resource_path = resource_path[6:]
        if resource_path and not (project_root / resource_path).is_file():
            fail(sheet.title, f"E{row}", f"resource does not exist: {resource_path}")
        entries[code] = LegendEntry(
            code=code,
            resource_path=resource_path,
            footprint_width=footprint_width,
            footprint_height=footprint_height,
            passability=clean_text(cells[16].value),
        )
    missing = (BASE_CODES | OBJECT_CODES) - entries.keys()
    if missing:
        raise ConversionError(
            "Legend is missing required codes: " + ", ".join(sorted(missing))
        )
    return entries


def parse_enemy_configs(project_root: Path) -> dict[int, dict[str, str]]:
    workbook_path = project_root / "design_tables" / "xlsx" / "actors.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "enemies" not in workbook.sheetnames:
        raise ConversionError("actors.xlsx is missing the enemies sheet")
    sheet = workbook["enemies"]
    headers = [clean_text(cell.value) for cell in sheet[1]]
    try:
        id_column = headers.index("enemy_id") + 1
        type_column = headers.index("enemy_type") + 1
        resource_column = headers.index("resource_type") + 1
    except ValueError as error:
        raise ConversionError(
            "actors.xlsx/enemies must contain enemy_id, enemy_type, and resource_type"
        ) from error
    configs: dict[int, dict[str, str]] = {}
    blank_rows = 0
    rows = sheet.iter_rows(min_row=4, max_row=1000, values_only=True)
    for row, values in enumerate(rows, start=4):
        raw_id = values[id_column - 1] if id_column <= len(values) else None
        if raw_id in (None, ""):
            blank_rows += 1
            if blank_rows >= 50:
                break
            continue
        blank_rows = 0
        enemy_id = int(raw_id)
        configs[enemy_id] = {
            "enemy_type": clean_text(values[type_column - 1]),
            "resource_type": clean_text(values[resource_column - 1]),
        }
    return configs


def parse_loot_group_ids(project_root: Path) -> set[int]:
    workbook_path = project_root / "design_tables" / "xlsx" / "loot.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "loot_groups" not in workbook.sheetnames:
        raise ConversionError("loot.xlsx is missing the loot_groups sheet")
    sheet = workbook["loot_groups"]
    headers = [clean_text(cell.value) for cell in sheet[1]]
    try:
        id_column = headers.index("loot_group_id")
    except ValueError as error:
        raise ConversionError(
            "loot.xlsx/loot_groups must contain loot_group_id"
        ) from error
    ids: set[int] = set()
    blank_rows = 0
    for values in sheet.iter_rows(
        min_row=4, max_row=1000, values_only=True
    ):
        raw_id = values[id_column] if id_column < len(values) else None
        if raw_id in (None, ""):
            blank_rows += 1
            if blank_rows >= 50:
                break
            continue
        blank_rows = 0
        ids.add(int(raw_id))
    return ids


def parse_object_code(token: str) -> tuple[str, str, bool]:
    anchored = token.endswith("@")
    normalized = token[:-1] if anchored else token
    code = normalized.split(":", 1)[0]
    return code, normalized, anchored


def token_parts(token: str) -> list[str]:
    return [part.strip() for part in token.split(":")]


def parse_map(
    sheet: Any,
    legend: dict[str, LegendEntry],
    enemy_configs: dict[int, dict[str, str]],
    loot_group_ids: set[int],
) -> MapSpec:
    map_id = clean_text(sheet["B3"].value)
    map_name = clean_text(sheet["D3"].value)
    width = parse_positive_int(sheet.title, "F3", sheet["F3"].value)
    height = parse_positive_int(sheet.title, "H3", sheet["H3"].value)
    cell_size = parse_positive_int(sheet.title, "J3", sheet["J3"].value)
    default_terrain = clean_text(sheet["L3"].value)
    output_scene = clean_text(sheet["G4"].value)
    is_default = parse_default_flag(sheet.title, sheet["L4"].value)

    if not re.fullmatch(r"[a-z0-9_]+", map_id):
        fail(sheet.title, "B3", "map_id must contain only lowercase letters, numbers, _")
    if sheet.title != f"map_{map_id}":
        fail(sheet.title, "B3", f"sheet name must be map_{map_id}")
    if not map_name:
        fail(sheet.title, "D3", "map_name cannot be empty")
    if cell_size != 32:
        fail(sheet.title, "J3", "cell_size must currently be 32")
    if default_terrain not in BASE_CODES:
        fail(sheet.title, "L3", f"unknown default terrain '{default_terrain}'")
    expected_scene = f"map_{map_id}.tscn"
    if output_scene != expected_scene:
        fail(sheet.title, "G4", f"output scene must be {expected_scene}")

    table_rows = list(
        sheet.iter_rows(
            min_row=GRID_START_ROW - 1,
            max_row=GRID_START_ROW + height - 1,
            min_col=GRID_START_COLUMN - 1,
            max_col=GRID_START_COLUMN + width - 1,
            values_only=True,
        )
    )
    header_row = table_rows[0]
    for x in range(width):
        value = header_row[1 + x]
        if value != x:
            fail(sheet.title, cell_name(x, -1), f"column header must be {x}")
    for y in range(height):
        value = table_rows[1 + y][0]
        if value != y:
            fail(sheet.title, f"A{GRID_START_ROW + y}", f"row header must be {y}")

    grid: list[list[str]] = []
    occurrences: dict[str, list[tuple[int, int, bool]]] = {}
    spawn_occurrences: list[Placement] = []
    for y in range(height):
        source_row = table_rows[1 + y]
        row: list[str] = []
        for x in range(width):
            raw_value = clean_text(source_row[1 + x])
            if not raw_value:
                fail(sheet.title, cell_name(x, y), "map cell cannot be empty")
            parts = [part.strip() for part in raw_value.split(";")]
            base = parts[0]
            if base not in BASE_CODES:
                fail(sheet.title, cell_name(x, y), f"unknown terrain code '{base}'")
            if (
                x in (0, width - 1) or y in (0, height - 1)
            ) and base != "X":
                fail(sheet.title, cell_name(x, y), "outer boundary terrain must be X")

            for object_token in parts[1:]:
                if not object_token:
                    fail(sheet.title, cell_name(x, y), "empty object token after ';'")
                code, normalized, anchored = parse_object_code(object_token)
                if code not in OBJECT_CODES:
                    fail(sheet.title, cell_name(x, y), f"unknown object code '{code}'")
                occurrences.setdefault(normalized, []).append((x, y, anchored))
                if code == "P":
                    spawn_occurrences.append(
                        Placement(code=code, token=normalized, x=x, y=y)
                    )
            object_codes_in_cell = {
                parse_object_code(token)[0] for token in parts[1:]
            }
            if "P" in object_codes_in_cell:
                if base not in {"G", "R", "BR"}:
                    fail(sheet.title, cell_name(x, y), "player spawn terrain is blocked")
                blocked_objects = object_codes_in_cell & {"T", "H1", "H2", "RB"}
                if blocked_objects:
                    fail(
                        sheet.title,
                        cell_name(x, y),
                        "player spawn overlaps blocked object(s): "
                        + ", ".join(sorted(blocked_objects)),
                    )
            row.append(raw_value)
        grid.append(row)

    if len(spawn_occurrences) != 1:
        raise ConversionError(
            f"[{sheet.title}] expected exactly one player spawn, found "
            f"{len(spawn_occurrences)}"
        )

    placements: list[Placement] = []
    for normalized, cells in occurrences.items():
        code = normalized.split(":", 1)[0]
        if code in {"N", "E", "BO"}:
            for x, y, _anchored in cells:
                validate_single_cell_token(
                    sheet.title,
                    x,
                    y,
                    normalized,
                    enemy_configs,
                    loot_group_ids,
                )
            anchors = [(x, y) for x, y, anchored in cells if anchored]
            if not anchors:
                for x, y, _anchored in cells:
                    placements.append(
                        Placement(code=code, token=normalized, x=x, y=y)
                    )
                continue

            enemy_id = int(token_parts(normalized)[1])
            resource_type = enemy_configs[enemy_id]["resource_type"]
            source_size = ENEMY_RESOURCE_SIZES.get(resource_type, (16, 16))
            footprint_width = max(1, math.ceil(source_size[0] * 2 / cell_size))
            footprint_height = max(1, math.ceil(source_size[1] * 2 / cell_size))
            cell_set = {(x, y) for x, y, _anchored in cells}
            covered: set[tuple[int, int]] = set()
            for anchor_x, anchor_y in anchors:
                for offset_y in range(footprint_height):
                    for offset_x in range(footprint_width):
                        point = (anchor_x + offset_x, anchor_y + offset_y)
                        if point not in cell_set:
                            fail(
                                sheet.title,
                                cell_name(*point),
                                f"missing legacy {code} continuation for @ anchor",
                            )
                        covered.add(point)
                placements.append(
                    Placement(code=code, token=normalized, x=anchor_x, y=anchor_y)
                )
            uncovered = cell_set - covered
            if uncovered:
                first = sorted(uncovered, key=lambda point: (point[1], point[0]))[0]
                fail(
                    sheet.title,
                    cell_name(*first),
                    f"legacy {code} continuation has no matching @ anchor",
                )
            continue
        if code in SINGLE_CELL_CODES:
            for x, y, _anchored in cells:
                validate_single_cell_token(
                    sheet.title,
                    x,
                    y,
                    normalized,
                    enemy_configs,
                    loot_group_ids,
                )
                placements.append(Placement(code=code, token=normalized, x=x, y=y))
            continue

        entry = legend[code]
        footprint_width = entry.footprint_width
        footprint_height = entry.footprint_height
        anchors = [(x, y) for x, y, anchored in cells if anchored]
        if footprint_width == 1 and footprint_height == 1:
            for x, y, _anchored in cells:
                placements.append(Placement(code=code, token=normalized, x=x, y=y))
            continue
        if not anchors:
            first_x, first_y, _ = cells[0]
            fail(
                sheet.title,
                cell_name(first_x, first_y),
                f"{code} occupies {footprint_width}x{footprint_height}; top-left cell needs @",
            )

        cell_set = {(x, y) for x, y, _anchored in cells}
        covered: set[tuple[int, int]] = set()
        for anchor_x, anchor_y in anchors:
            if anchor_x + footprint_width > width or anchor_y + footprint_height > height:
                fail(sheet.title, cell_name(anchor_x, anchor_y), f"{code} footprint is out of bounds")
            for offset_y in range(footprint_height):
                for offset_x in range(footprint_width):
                    point = (anchor_x + offset_x, anchor_y + offset_y)
                    if point not in cell_set:
                        fail(
                            sheet.title,
                            cell_name(*point),
                            f"missing {code} continuation for anchor at {cell_name(anchor_x, anchor_y)}",
                        )
                    if point in covered:
                        fail(sheet.title, cell_name(*point), f"overlapping {code} footprints")
                    covered.add(point)
            placements.append(
                Placement(
                    code=code,
                    token=normalized,
                    x=anchor_x,
                    y=anchor_y,
                    width=footprint_width,
                    height=footprint_height,
                )
            )
        uncovered = cell_set - covered
        if uncovered:
            first = sorted(uncovered, key=lambda point: (point[1], point[0]))[0]
            fail(sheet.title, cell_name(*first), f"{code} continuation has no matching @ anchor")

    return MapSpec(
        sheet_name=sheet.title,
        map_id=map_id,
        map_name=map_name,
        width=width,
        height=height,
        cell_size=cell_size,
        default_terrain=default_terrain,
        output_scene=output_scene,
        is_default=is_default,
        grid=grid,
        placements=placements,
        spawn=spawn_occurrences[0],
    )


def validate_single_cell_token(
    sheet: str,
    x: int,
    y: int,
    token: str,
    enemy_configs: dict[int, dict[str, str]],
    loot_group_ids: set[int],
) -> None:
    parts = token_parts(token)
    code = parts[0]
    if code == "P":
        if len(parts) > 2 or (len(parts) == 2 and parts[1] not in {"U", "D", "L", "R"}):
            fail(sheet, cell_name(x, y), "player syntax is P or P:U/D/L/R")
        return
    if code == "C":
        if len(parts) != 2 or not parts[1].isdigit() or int(parts[1]) <= 0:
            fail(sheet, cell_name(x, y), "chest syntax is C:<positive loot_group_id>")
        if int(parts[1]) not in loot_group_ids:
            fail(sheet, cell_name(x, y), f"loot_group_id {parts[1]} does not exist")
        return
    if code not in {"N", "E", "BO"}:
        if len(parts) != 1:
            fail(sheet, cell_name(x, y), f"{code} does not accept parameters")
        return
    if len(parts) < 2 or not parts[1].isdigit():
        fail(sheet, cell_name(x, y), f"{code} requires a numeric enemy_id")
    enemy_id = int(parts[1])
    if enemy_id not in enemy_configs:
        fail(sheet, cell_name(x, y), f"enemy_id {enemy_id} is not in actors.xlsx/enemies")
    enemy_type = enemy_configs[enemy_id]["enemy_type"]
    if code == "BO" and enemy_type != "boss":
        fail(sheet, cell_name(x, y), f"BO requires a boss enemy, got {enemy_type}")
    if code in {"N", "E"} and enemy_type == "boss":
        fail(sheet, cell_name(x, y), f"{code} cannot reference a boss enemy")
    if len(parts) >= 3 and parts[2] not in {"LR", "RL", "UD", "DU", "NONE"}:
        fail(sheet, cell_name(x, y), "patrol direction must be LR/RL/UD/DU/NONE")
    if len(parts) >= 4:
        try:
            if float(parts[3]) < 0:
                raise ValueError
        except ValueError:
            fail(sheet, cell_name(x, y), "patrol radius must be a non-negative number")
    if len(parts) >= 5:
        try:
            if float(parts[4]) < 0:
                raise ValueError
        except ValueError:
            fail(sheet, cell_name(x, y), "patrol speed must be a non-negative number")
    if len(parts) > 5:
        fail(sheet, cell_name(x, y), "enemy token has too many parameters")


def resource_digest(project_root: Path, entry: LegendEntry) -> str:
    if not entry.resource_path:
        return ""
    return hashlib.sha256((project_root / entry.resource_path).read_bytes()).hexdigest()


def map_payload(
    spec: MapSpec,
    legend: dict[str, LegendEntry],
    project_root: Path,
) -> dict[str, Any]:
    used_codes = {cell.split(";", 1)[0] for row in spec.grid for cell in row}
    used_codes.update(placement.code for placement in spec.placements)
    resources = {
        code: {
            "path": legend[code].resource_path,
            "sha256": resource_digest(project_root, legend[code]),
            "footprint": [legend[code].footprint_width, legend[code].footprint_height],
        }
        for code in sorted(used_codes)
    }
    return {
        "tool_version": TOOL_VERSION,
        "map_id": spec.map_id,
        "map_name": spec.map_name,
        "width": spec.width,
        "height": spec.height,
        "cell_size": spec.cell_size,
        "default_terrain": spec.default_terrain,
        "grid": spec.grid,
        "placements": [placement.__dict__ for placement in spec.placements],
        "resources": resources,
    }


def payload_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def world_position(spec: MapSpec, x: int, y: int, width: int = 1, height: int = 1) -> tuple[float, float]:
    half_width = spec.width * spec.cell_size / 2.0
    half_height = spec.height * spec.cell_size / 2.0
    return (
        (x + width / 2.0) * spec.cell_size - half_width,
        (y + height / 2.0) * spec.cell_size - half_height,
    )


def tscn_string(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def build_generated_scene(
    spec: MapSpec,
    enemy_configs: dict[int, dict[str, str]],
) -> str:
    lines = [
        "[gd_scene load_steps=4 format=3]",
        "",
        '[ext_resource type="Script" path="res://scripts/generated_level.gd" id="1_level"]',
        '[ext_resource type="PackedScene" path="res://scenes/enemy.tscn" id="2_enemy"]',
        '[ext_resource type="PackedScene" path="res://scenes/chest.tscn" id="3_chest"]',
        "",
        f'[node name="Map{camel_name(spec.map_id)}Generated" type="Node2D" groups=["generated_level"]]',
        "y_sort_enabled = true",
        'script = ExtResource("1_level")',
        f"map_id = {tscn_string(spec.map_id)}",
        f"map_name = {tscn_string(spec.map_name)}",
        f"map_width = {spec.width}",
        f"map_height = {spec.height}",
        f"cell_size = {spec.cell_size}",
        f'data_path = {tscn_string(f"res://data/generated/levels/map_{spec.map_id}.json")}',
        "",
        '[node name="Terrain" type="Node2D" parent="."]',
        "z_index = -100",
        "",
        '[node name="GroundLayer" type="Node2D" parent="Terrain"]',
        '[node name="RoadLayer" type="Node2D" parent="Terrain"]',
        '[node name="WaterLayer" type="Node2D" parent="Terrain"]',
        '[node name="BridgeLayer" type="Node2D" parent="Terrain"]',
        "",
        '[node name="Environment" type="Node2D" parent="."]',
        "y_sort_enabled = true",
        '[node name="Trees" type="Node2D" parent="Environment"]',
        '[node name="Bushes" type="Node2D" parent="Environment"]',
        '[node name="GrassDecorations" type="Node2D" parent="Environment"]',
        '[node name="Houses" type="Node2D" parent="Environment"]',
        '[node name="RegionBoundaries" type="Node2D" parent="Environment"]',
        "",
        '[node name="Gameplay" type="Node2D" parent="."]',
        "y_sort_enabled = true",
        '[node name="PlayerSpawn" type="Marker2D" parent="Gameplay" groups=["player_spawn"]]',
    ]
    spawn_x, spawn_y = world_position(spec, spec.spawn.x, spec.spawn.y)
    spawn_parts = token_parts(spec.spawn.token)
    facing = spawn_parts[1] if len(spawn_parts) > 1 else "D"
    lines.extend(
        [
            f"position = Vector2({format_number(spawn_x)}, {format_number(spawn_y)})",
            f"metadata/facing = {tscn_string(facing)}",
            "",
            '[node name="Chests" type="Node2D" parent="Gameplay"]',
            '[node name="NormalEnemies" type="Node2D" parent="Gameplay"]',
            '[node name="EliteEnemies" type="Node2D" parent="Gameplay"]',
            '[node name="Bosses" type="Node2D" parent="Gameplay"]',
            "",
            '[node name="Collisions" type="Node2D" parent="."]',
            '[node name="MapBoundary" type="Node2D" parent="Collisions"]',
            '[node name="WaterCollision" type="Node2D" parent="Collisions"]',
            '[node name="PropCollision" type="Node2D" parent="Collisions"]',
            '[node name="RegionBoundaryCollision" type="Node2D" parent="Collisions"]',
        ]
    )

    chest_index = 0
    enemy_index = 0
    for placement in spec.placements:
        if placement.code == "C":
            chest_index += 1
            parts = token_parts(placement.token)
            loot_group_id = int(parts[1])
            x, y = world_position(spec, placement.x, placement.y)
            lines.extend(
                [
                    "",
                    f'[node name="Chest_{chest_index}_{placement.x}_{placement.y}" parent="Gameplay/Chests" instance=ExtResource("3_chest")]',
                    f"position = Vector2({format_number(x)}, {format_number(y)})",
                    f'container_instance_id = {tscn_string(f"{spec.map_id}_chest_{placement.x}_{placement.y}")}',
                    f"loot_group_id = {loot_group_id}",
                ]
            )
        elif placement.code in {"N", "E", "BO"}:
            enemy_index += 1
            parts = token_parts(placement.token)
            enemy_id = int(parts[1])
            direction = parts[2] if len(parts) >= 3 else "LR"
            radius_cells = float(parts[3]) if len(parts) >= 4 else 2.5
            speed = float(parts[4]) if len(parts) >= 5 else 70.0
            resource_type = enemy_configs[enemy_id]["resource_type"]
            parent = {
                "N": "NormalEnemies",
                "E": "EliteEnemies",
                "BO": "Bosses",
            }[placement.code]
            x, y = world_position(spec, placement.x, placement.y)
            lines.extend(
                [
                    "",
                    f'[node name="{placement.code}_Enemy_{enemy_index}_{placement.x}_{placement.y}" parent="Gameplay/{parent}" instance=ExtResource("2_enemy")]',
                    f"position = Vector2({format_number(x)}, {format_number(y)})",
                    f"enemy_id = {enemy_id}",
                    f"animal_type = {tscn_string(resource_type)}",
                    f"patrol_direction = {tscn_string(direction)}",
                    f"patrol_distance = {format_number(radius_cells * spec.cell_size)}",
                    f"patrol_speed = {format_number(speed)}",
                    f'map_object_id = {tscn_string(f"{spec.map_id}_enemy_{placement.x}_{placement.y}")}',
                ]
            )

    return "\n".join(lines) + "\n"


def build_wrapper_scene(spec: MapSpec) -> str:
    return "\n".join(
        [
            "[gd_scene load_steps=2 format=3]",
            "",
            f'[ext_resource type="PackedScene" path="res://scenes/maps/generated/map_{spec.map_id}_generated.tscn" id="1_generated"]',
            "",
            f'[node name="Map{camel_name(spec.map_id)}" type="Node2D"]',
            "y_sort_enabled = true",
            "",
            '[node name="Generated" parent="." instance=ExtResource("1_generated")]',
            "",
            '[node name="ManualOverrides" type="Node2D" parent="."]',
            "y_sort_enabled = true",
            "",
        ]
    )


def camel_name(value: str) -> str:
    return "".join(part.capitalize() for part in value.split("_"))


def format_number(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.4f}".rstrip("0").rstrip(".")


def atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", newline="\n", dir=path.parent, delete=False
    ) as handle:
        handle.write(content)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def atomic_write_json(path: Path, data: Any) -> None:
    atomic_write_text(
        path,
        json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
    )


def load_manifest(source_root: Path, output_root: Path) -> dict[str, Any]:
    relative = Path("data/generated/levels/manifest.json")
    for root in (output_root, source_root):
        path = root / relative
        if path.is_file():
            try:
                return json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                return {}
    return {}


def changed_cells(old_grid: Any, new_grid: list[list[str]]) -> list[str]:
    if not isinstance(old_grid, list):
        return []
    changes = []
    max_height = max(len(old_grid), len(new_grid))
    for y in range(max_height):
        old_row = old_grid[y] if y < len(old_grid) and isinstance(old_grid[y], list) else []
        new_row = new_grid[y] if y < len(new_grid) else []
        max_width = max(len(old_row), len(new_row))
        for x in range(max_width):
            old_value = old_row[x] if x < len(old_row) else None
            new_value = new_row[x] if x < len(new_row) else None
            if old_value != new_value:
                changes.append(f"({x},{y}) {old_value!r} -> {new_value!r}")
    return changes


def update_world_scene(source_root: Path, output_root: Path, default_map_id: str) -> bool:
    relative = Path("scenes/world.tscn")
    output_path = output_root / relative
    source_path = output_path if output_path.is_file() else source_root / relative
    content = source_path.read_text(encoding="utf-8")
    pattern = re.compile(
        r'(\[ext_resource type="PackedScene" path=")[^"]+(" id="4_map"\])'
    )
    replacement_path = f"res://scenes/maps/map_{default_map_id}.tscn"
    updated, count = pattern.subn(rf"\g<1>{replacement_path}\g<2>", content, count=1)
    if count != 1:
        raise ConversionError(
            "scenes/world.tscn must contain the map ext_resource with id=4_map"
        )
    if updated == content and output_path.is_file():
        return False
    atomic_write_text(output_path, updated)
    return updated != content


def main() -> None:
    args = parse_args()
    project_root = args.project_root.resolve()
    output_root = (args.output_root or project_root).resolve()
    workbook_path = (
        args.workbook.resolve()
        if args.workbook
        else project_root / "design_tables" / "xlsx" / "levels.xlsx"
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    legend = parse_legend(workbook, project_root)
    enemy_configs = parse_enemy_configs(project_root)
    loot_group_ids = parse_loot_group_ids(project_root)

    map_sheets = [
        workbook[name]
        for name in workbook.sheetnames
        if name.startswith(MAP_SHEET_PREFIX)
    ]
    if not map_sheets:
        raise ConversionError("levels.xlsx contains no map_ sheets")

    specs = [
        parse_map(sheet, legend, enemy_configs, loot_group_ids)
        for sheet in map_sheets
    ]
    ids = [spec.map_id for spec in specs]
    duplicates = sorted({map_id for map_id in ids if ids.count(map_id) > 1})
    if duplicates:
        raise ConversionError("Duplicate map_id values: " + ", ".join(duplicates))
    defaults = [spec for spec in specs if spec.is_default]
    if len(defaults) != 1:
        raise ConversionError(
            f"Expected exactly one default map, found {len(defaults)}"
        )

    payloads = {
        spec.map_id: map_payload(spec, legend, project_root) for spec in specs
    }
    hashes = {map_id: payload_hash(payload) for map_id, payload in payloads.items()}
    previous_manifest = load_manifest(project_root, output_root)
    previous_maps = previous_manifest.get("maps", {})
    next_manifest: dict[str, Any] = {
        "tool_version": TOOL_VERSION,
        "default_map_id": defaults[0].map_id,
        "maps": {},
    }

    changed_count = 0
    for spec in specs:
        previous = previous_maps.get(spec.map_id, {})
        changed = previous.get("hash") != hashes[spec.map_id]
        next_manifest["maps"][spec.map_id] = {
            "hash": hashes[spec.map_id],
            "sheet": spec.sheet_name,
            "grid": spec.grid,
        }
        if changed:
            changed_count += 1
            differences = changed_cells(previous.get("grid"), spec.grid)
            if previous:
                print(f"Rebuild {spec.map_id}: {len(differences)} changed cells")
                for difference in differences[:50]:
                    print(f"  {difference}")
                if len(differences) > 50:
                    print(f"  ... {len(differences) - 50} more")
                if not differences:
                    print("  map metadata, legend, or a referenced resource changed")
            else:
                print(f"Build new map {spec.map_id}")
            atomic_write_json(
                output_root / f"data/generated/levels/map_{spec.map_id}.json",
                payloads[spec.map_id],
            )
            atomic_write_text(
                output_root
                / f"scenes/maps/generated/map_{spec.map_id}_generated.tscn",
                build_generated_scene(spec, enemy_configs),
            )
        else:
            print(f"Skip unchanged map {spec.map_id}")

        wrapper_relative = Path(f"scenes/maps/map_{spec.map_id}.tscn")
        wrapper_output = output_root / wrapper_relative
        wrapper_source = project_root / wrapper_relative
        if not wrapper_output.exists() and not wrapper_source.exists():
            atomic_write_text(wrapper_output, build_wrapper_scene(spec))
            print(f"Create wrapper {wrapper_relative.as_posix()}")

    world_changed = update_world_scene(project_root, output_root, defaults[0].map_id)
    atomic_write_json(
        output_root / "data/generated/levels/manifest.json", next_manifest
    )
    print(
        f"Level conversion complete: {changed_count} rebuilt, "
        f"{len(specs) - changed_count} unchanged; default={defaults[0].map_id}"
    )
    if world_changed:
        print("Updated scenes/world.tscn default map reference")


if __name__ == "__main__":
    try:
        main()
    except ConversionError as error:
        print(f"Level conversion failed: {error}", file=sys.stderr)
        raise SystemExit(1)
