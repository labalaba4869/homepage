from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "Missing dependency: openpyxl\n"
        "Run: py -3 -m pip install -r design_tables/requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1)


TOOL_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TOOL_DIR.parent
SOURCE_DIR = TOOL_DIR / "xlsx"
OUTPUT_DIR = PROJECT_ROOT / "data" / "generated" / "tables"
PARAM_SEPARATOR = ";"


@dataclass(frozen=True)
class TableSpec:
    workbook: str
    sheet: str
    output: str
    headers: tuple[str, ...]
    types: tuple[str, ...]
    id_field: str
    required_fields: tuple[str, ...] = ()
    int_fields: tuple[str, ...] = ()
    optional_int_fields: tuple[str, ...] = ()
    float_fields: tuple[str, ...] = ()
    int_array_fields: tuple[str, ...] = ()
    enum_fields: tuple[tuple[str, tuple[str, ...]], ...] = ()


TABLES = (
    TableSpec(
        workbook="cards.xlsx",
        sheet="cards",
        output="cards.csv",
        headers=(
            "card_id",
            "card_type",
            "card_name",
            "card_race",
            "card_equality",
            "card_attack",
            "card_maxhp",
            "cost",
            "card_description",
            "description_params",
        ),
        types=(
            "int",
            "string",
            "string",
            "string",
            "int",
            "int",
            "int",
            "int",
            "string",
            "int[]",
        ),
        id_field="card_id",
        required_fields=("card_type", "card_name"),
        int_fields=("card_equality",),
        optional_int_fields=("card_attack", "card_maxhp", "cost"),
        int_array_fields=("description_params",),
        enum_fields=(("card_type", ("function", "battle_hero")),),
    ),
    TableSpec(
        workbook="cards.xlsx",
        sheet="card_equality",
        output="card_equality.csv",
        headers=("equality_id", "equality_name", "use_num"),
        types=("int", "string", "int"),
        id_field="equality_id",
        required_fields=("equality_name",),
        int_fields=("use_num",),
    ),
    TableSpec(
        workbook="cards.xlsx",
        sheet="deck",
        output="decks.csv",
        headers=("deck_id", "card_type"),
        types=("int", "int[]"),
        id_field="deck_id",
        int_array_fields=("card_type",),
    ),
    TableSpec(
        workbook="actors.xlsx",
        sheet="player_npc",
        output="player_npc.csv",
        headers=(
            "actor_id",
            "actor_type",
            "name",
            "deck_id",
            "max_hp",
            "initial_energy",
            "move_speed",
        ),
        types=("int", "string", "string", "int", "int", "int", "float"),
        id_field="actor_id",
        required_fields=("actor_type", "name"),
        int_fields=("deck_id", "max_hp", "initial_energy"),
        float_fields=("move_speed",),
        enum_fields=(("actor_type", ("player", "npc")),),
    ),
    TableSpec(
        workbook="actors.xlsx",
        sheet="enemies",
        output="enemies.csv",
        headers=(
            "enemy_id",
            "enemy_type",
            "name",
            "deck_id",
            "max_hp",
            "initial_energy",
            "resource_type",
        ),
        types=("int", "string", "string", "int", "int", "int", "string"),
        id_field="enemy_id",
        required_fields=("enemy_type", "name", "resource_type"),
        int_fields=("deck_id", "max_hp", "initial_energy"),
        enum_fields=(("enemy_type", ("normal", "boss")),),
    ),
)


def normalize(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip().replace("\r\n", "\\n").replace("\n", "\\n")
    return value


def parse_integer(value: Any, location: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{location} must be an integer")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value.strip()):
        return int(value)
    raise ValueError(f"{location} must be an integer")


def parse_int_array(value: Any, location: str) -> None:
    if value is None or str(value).strip() == "":
        return
    for index, item in enumerate(str(value).split(PARAM_SEPARATOR)):
        if item.strip() == "":
            raise ValueError(f"{location}[{index}] cannot be empty")
        parse_integer(item.strip(), f"{location}[{index}]")


def validate_record(
    spec: TableSpec,
    record: dict[str, Any],
    location: str,
) -> None:
    for field in spec.required_fields:
        if str(record[field]).strip() == "":
            raise ValueError(f"{location}/{field} cannot be empty")
    for field in spec.int_fields:
        parse_integer(record[field], f"{location}/{field}")
    for field in spec.optional_int_fields:
        if str(record[field]).strip() != "":
            parse_integer(record[field], f"{location}/{field}")
    for field in spec.float_fields:
        try:
            float(record[field])
        except (TypeError, ValueError) as error:
            raise ValueError(f"{location}/{field} must be a number") from error
    for field in spec.int_array_fields:
        parse_int_array(record[field], f"{location}/{field}")
    for field, allowed_values in spec.enum_fields:
        if str(record[field]).strip() not in allowed_values:
            raise ValueError(
                f"{location}/{field} must be one of {allowed_values}"
            )


def read_table(spec: TableSpec) -> list[list[Any]]:
    workbook_path = SOURCE_DIR / spec.workbook
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing workbook: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if spec.sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{spec.workbook} is missing sheet '{spec.sheet}'")

    sheet = workbook[spec.sheet]
    raw_rows = [
        list(row[: len(spec.headers)])
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    if len(raw_rows) < 3:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} requires header, type and note rows"
        )

    actual_headers = tuple(str(value or "").strip() for value in raw_rows[0])
    actual_types = tuple(str(value or "").strip() for value in raw_rows[1])
    if actual_headers != spec.headers:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} headers changed.\n"
            f"Expected: {spec.headers}\nActual:   {actual_headers}"
        )
    if actual_types != spec.types:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} types changed.\n"
            f"Expected: {spec.types}\nActual:   {actual_types}"
        )

    rows = [raw_rows[0], raw_rows[1], raw_rows[2]]
    seen_ids: set[int] = set()
    for row_number, raw_row in enumerate(raw_rows[3:], start=4):
        if all(value is None or str(value).strip() == "" for value in raw_row):
            continue
        row = [normalize(value) for value in raw_row]
        record = dict(zip(spec.headers, row, strict=True))
        location = f"{spec.workbook}/{spec.sheet} row {row_number}"
        record_id = parse_integer(
            record[spec.id_field],
            f"{location}/{spec.id_field}",
        )
        if record_id in seen_ids:
            raise ValueError(
                f"{location} has duplicate {spec.id_field} {record_id}"
            )
        seen_ids.add(record_id)
        validate_record(spec, record, location)
        rows.append(row)
    return rows


def write_csv(output_path: Path, rows: list[list[Any]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".csv.tmp")
    with temporary_path.open("w", encoding="utf-8", newline="") as file:
        csv.writer(file, lineterminator="\n").writerows(rows)
    temporary_path.replace(output_path)


def main() -> None:
    converted: list[tuple[str, int]] = []
    for spec in TABLES:
        rows = read_table(spec)
        write_csv(OUTPUT_DIR / spec.output, rows)
        converted.append((spec.output, max(0, len(rows) - 3)))
    print("Table conversion complete:")
    for filename, data_rows in converted:
        print(f"  {filename}: {data_rows} data rows")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as error:
        print(f"Table conversion failed: {error}", file=sys.stderr)
        raise SystemExit(1)
