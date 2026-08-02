from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "Missing dependency: openpyxl\n"
        "Run: py -3 -m pip install -r design_tables/requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1)


TOOL_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TOOL_DIR.parent
SOURCE_DIR = TOOL_DIR / "xlsx"
OUTPUT_DIR = PROJECT_ROOT / "data" / "generated" / "tables"
PARAM_SEPARATOR = ";"
LEGACY_OUTPUTS = (
    "loot_entries.csv",
    "search_containers.csv",
)
MACHINE_VALUE_MAPPINGS: dict[tuple[str, str], dict[str, str]] = {}


@dataclass(frozen=True)
class TableSpec:
    workbook: str
    sheet: str
    output: str
    headers: tuple[str, ...]
    types: tuple[str, ...]
    id_field: str
    required_fields: tuple[str, ...] = ()
    int_fields: tuple[str, ...] = ()
    optional_int_fields: tuple[str, ...] = ()
    float_fields: tuple[str, ...] = ()
    int_array_fields: tuple[str, ...] = ()
    enum_fields: tuple[tuple[str, tuple[str, ...]], ...] = ()


TABLES = (
    TableSpec(
        workbook="cards.xlsx",
        sheet="cards",
        output="cards.csv",
        headers=(
            "card_id",
            "card_type",
            "card_name",
            "card_race",
            "card_equality",
            "card_attack",
            "card_maxhp",
            "cost",
            "card_description",
            "description_params",
            "card_drop_weight",
        ),
        types=(
            "int",
            "string",
            "string",
            "string",
            "int",
            "int",
            "int",
            "int",
            "string",
            "int[]",
            "float",
        ),
        id_field="card_id",
        required_fields=("card_type", "card_name"),
        int_fields=("card_equality",),
        optional_int_fields=("card_attack", "card_maxhp", "cost"),
        float_fields=("card_drop_weight",),
        int_array_fields=("description_params",),
        enum_fields=(("card_type", ("function", "battle_hero")),),
    ),
    TableSpec(
        workbook="cards.xlsx",
        sheet="card_equality",
        output="card_equality.csv",
        headers=("equality_id", "equality_name", "use_num"),
        types=("int", "string", "int"),
        id_field="equality_id",
        required_fields=("equality_name",),
        int_fields=("use_num",),
    ),
    TableSpec(
        workbook="cards.xlsx",
        sheet="deck",
        output="decks.csv",
        headers=("deck_id", "card_type"),
        types=("int", "int[]"),
        id_field="deck_id",
        int_array_fields=("card_type",),
    ),
    TableSpec(
        workbook="actors.xlsx",
        sheet="player_npc",
        output="player_npc.csv",
        headers=(
            "actor_id",
            "actor_type",
            "name",
            "deck_id",
            "max_hp",
            "initial_energy",
            "move_speed",
            "inventory_slots",
            "max_carry_weight",
            "initial_item_ids",
            "initial_item_counts",
        ),
        types=(
            "int",
            "string",
            "string",
            "int",
            "int",
            "int",
            "float",
            "int",
            "float",
            "int[]",
            "int[]",
        ),
        id_field="actor_id",
        required_fields=("actor_type", "name"),
        int_fields=(
            "deck_id",
            "max_hp",
            "initial_energy",
            "inventory_slots",
        ),
        float_fields=("move_speed", "max_carry_weight"),
        int_array_fields=("initial_item_ids", "initial_item_counts"),
        enum_fields=(("actor_type", ("player", "npc")),),
    ),
    TableSpec(
        workbook="actors.xlsx",
        sheet="enemies",
        output="enemies.csv",
        headers=(
            "enemy_id",
            "enemy_type",
            "name",
            "deck_id",
            "max_hp",
            "initial_energy",
            "resource_type",
        ),
        types=("int", "string", "string", "int", "int", "int", "string"),
        id_field="enemy_id",
        required_fields=("enemy_type", "name", "resource_type"),
        int_fields=("deck_id", "max_hp", "initial_energy"),
        enum_fields=(("enemy_type", ("normal", "boss")),),
    ),
    TableSpec(
        workbook="loot.xlsx",
        sheet="items",
        output="items.csv",
        headers=(
            "item_id",
            "item_name",
            "item_equality",
            "drop_weight",
            "max_stack",
            "sell_price",
            "unit_weight",
            "grid_width",
            "grid_height",
            "icon_path",
            "item_description",
        ),
        types=(
            "int",
            "string",
            "int",
            "float",
            "int",
            "int",
            "float",
            "int",
            "int",
            "string",
            "string",
        ),
        id_field="item_id",
        required_fields=("item_name",),
        int_fields=(
            "item_equality",
            "max_stack",
            "sell_price",
            "grid_width",
            "grid_height",
        ),
        float_fields=("drop_weight", "unit_weight"),
    ),
    TableSpec(
        workbook="loot.xlsx",
        sheet="active_items",
        output="active_items.csv",
        headers=(
            "active_item_id",
            "active_item_name",
            "item_equality",
            "unit_weight",
            "item_description",
            "effect_params",
            "sell_price",
            "drop_weight",
            "max_carry_count",
            "grid_width",
            "grid_height",
            "icon_path",
        ),
        types=(
            "int",
            "string",
            "int",
            "float",
            "string",
            "int[]",
            "int",
            "float",
            "int",
            "int",
            "int",
            "string",
        ),
        id_field="active_item_id",
        required_fields=("active_item_name", "item_description"),
        int_fields=(
            "sell_price",
            "max_carry_count",
            "item_equality",
            "grid_width",
            "grid_height",
        ),
        float_fields=("unit_weight", "drop_weight"),
        int_array_fields=("effect_params",),
    ),
    TableSpec(
        workbook="loot.xlsx",
        sheet="item_equality",
        output="item_equality.csv",
        headers=(
            "equality_id",
            "equality_name",
            "search_time",
            "frame_path",
        ),
        types=("int", "string", "float", "string"),
        id_field="equality_id",
        required_fields=("equality_name", "frame_path"),
        float_fields=("search_time",),
    ),
    TableSpec(
        workbook="loot.xlsx",
        sheet="loot_groups",
        output="loot_groups.csv",
        headers=(
            "loot_group_id",
            "loot_group_name",
            "slot_count",
            "fill_chance",
            "white_chance",
            "green_chance",
            "blue_chance",
            "purple_chance",
            "gold_chance",
            "normal_item_chance",
            "active_item_chance",
            "card_chance",
        ),
        types=(
            "int",
            "string",
            "int",
            "float",
            "float",
            "float",
            "float",
            "float",
            "float",
            "float",
            "float",
            "float",
        ),
        id_field="loot_group_id",
        required_fields=("loot_group_name",),
        int_fields=("slot_count",),
        float_fields=(
            "fill_chance",
            "white_chance",
            "green_chance",
            "blue_chance",
            "purple_chance",
            "gold_chance",
            "normal_item_chance",
            "active_item_chance",
            "card_chance",
        ),
    ),
)


def normalize(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip().replace("\r\n", "\\n").replace("\n", "\\n")
    return value


def parse_integer(value: Any, location: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{location} must be an integer")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value.strip()):
        return int(value)
    raise ValueError(f"{location} must be an integer")


def parse_int_array(value: Any, location: str) -> None:
    if value is None or str(value).strip() == "":
        return
    for index, item in enumerate(str(value).split(PARAM_SEPARATOR)):
        if item.strip() == "":
            raise ValueError(f"{location}[{index}] cannot be empty")
        parse_integer(item.strip(), f"{location}[{index}]")


def validate_record(
    spec: TableSpec,
    record: dict[str, Any],
    location: str,
) -> None:
    for field in spec.required_fields:
        if str(record[field]).strip() == "":
            raise ValueError(f"{location}/{field} cannot be empty")
    for field in spec.int_fields:
        parse_integer(record[field], f"{location}/{field}")
    for field in spec.optional_int_fields:
        if str(record[field]).strip() != "":
            parse_integer(record[field], f"{location}/{field}")
    for field in spec.float_fields:
        try:
            float(record[field])
        except (TypeError, ValueError) as error:
            raise ValueError(f"{location}/{field} must be a number") from error
    for field in spec.int_array_fields:
        parse_int_array(record[field], f"{location}/{field}")
    for field, allowed_values in spec.enum_fields:
        if str(record[field]).strip() not in allowed_values:
            raise ValueError(
                f"{location}/{field} must be one of {allowed_values}"
            )


def read_table(spec: TableSpec) -> list[list[Any]]:
    workbook_path = SOURCE_DIR / spec.workbook
    if not workbook_path.exists():
        raise FileNotFoundError(f"Missing workbook: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if spec.sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{spec.workbook} is missing sheet '{spec.sheet}'")

    sheet = workbook[spec.sheet]
    raw_rows: list[list[Any]] = []
    for source_row in sheet.iter_rows(values_only=True):
        row = list(source_row[: len(spec.headers)])
        if len(row) < len(spec.headers):
            row.extend([None] * (len(spec.headers) - len(row)))
        raw_rows.append(row)
    workbook.close()
    if len(raw_rows) < 3:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} requires header, type and note rows"
        )

    actual_headers = tuple(str(value or "").strip() for value in raw_rows[0])
    actual_types = tuple(str(value or "").strip() for value in raw_rows[1])
    if actual_headers != spec.headers:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} headers changed.\n"
            f"Expected: {spec.headers}\nActual:   {actual_headers}"
        )
    if actual_types != spec.types:
        raise ValueError(
            f"{spec.workbook}/{spec.sheet} types changed.\n"
            f"Expected: {spec.types}\nActual:   {actual_types}"
        )

    rows = [raw_rows[0], raw_rows[1], raw_rows[2]]
    seen_ids: set[int] = set()
    for row_number, raw_row in enumerate(raw_rows[3:], start=4):
        if all(value is None or str(value).strip() == "" for value in raw_row):
            continue
        row = [normalize(value) for value in raw_row]
        record = dict(zip(spec.headers, row, strict=True))
        location = f"{spec.workbook}/{spec.sheet} row {row_number}"
        record_id = parse_integer(
            record[spec.id_field],
            f"{location}/{spec.id_field}",
        )
        if record_id in seen_ids:
            raise ValueError(
                f"{location} has duplicate {spec.id_field} {record_id}"
            )
        seen_ids.add(record_id)
        validate_record(spec, record, location)
        for field in spec.headers:
            mapping = MACHINE_VALUE_MAPPINGS.get((spec.output, field))
            if mapping is None:
                continue
            source_value = str(record[field]).strip()
            row[spec.headers.index(field)] = mapping[source_value]
        rows.append(row)
    return rows


def write_csv(output_path: Path, rows: list[list[Any]]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".csv.tmp")
    with temporary_path.open("w", encoding="utf-8", newline="") as file:
        csv.writer(file, lineterminator="\n").writerows(rows)
    temporary_path.replace(output_path)


def rows_to_records(rows: list[list[Any]]) -> list[dict[str, Any]]:
    headers = [str(value) for value in rows[0]]
    return [dict(zip(headers, row, strict=True)) for row in rows[3:]]


def validate_card_tables(tables: dict[str, list[list[Any]]]) -> None:
    cards = rows_to_records(tables["cards.csv"])
    for card in cards:
        card_id = int(card["card_id"])
        if float(card["card_drop_weight"]) <= 0:
            raise ValueError(
                f"cards/card_id {card_id} card_drop_weight must be > 0"
            )


def validate_loot_tables(tables: dict[str, list[list[Any]]]) -> None:
    normal_items = {
        int(record["item_id"]): record
        for record in rows_to_records(tables["items.csv"])
    }
    active_items = {
        int(record["active_item_id"]): record
        for record in rows_to_records(tables["active_items.csv"])
    }
    duplicate_ids = set(normal_items).intersection(active_items)
    if duplicate_ids:
        raise ValueError(
            f"normal and active item IDs overlap: {sorted(duplicate_ids)}"
        )

    items = dict(normal_items)
    for active_item_id, active_item in active_items.items():
        merged = dict(active_item)
        merged["item_id"] = active_item_id
        merged["item_name"] = active_item["active_item_name"]
        merged["max_stack"] = active_item["max_carry_count"]
        items[active_item_id] = merged

    qualities = {
        int(record["equality_id"]): record
        for record in rows_to_records(tables["item_equality.csv"])
    }
    if set(qualities) != set(range(5)):
        raise ValueError("item_equality must define IDs 0, 1, 2, 3 and 4")
    for quality_id, quality in qualities.items():
        if float(quality["search_time"]) <= 0:
            raise ValueError(
                f"item quality {quality_id} search_time must be > 0"
            )

    groups = {
        int(record["loot_group_id"]): record
        for record in rows_to_records(tables["loot_groups.csv"])
    }
    for item_id, item in normal_items.items():
        if item_id <= 0 or item_id >= 10000:
            raise ValueError(
                f"normal item_id {item_id} must use the 1..9999 range"
            )
        quality = int(item["item_equality"])
        if quality not in qualities:
            raise ValueError(f"items/item_id {item_id} has invalid quality")
        if int(item["max_stack"]) <= 0:
            raise ValueError(f"items/item_id {item_id} max_stack must be > 0")
        if int(item["max_stack"]) > 3:
            raise ValueError(f"items/item_id {item_id} max_stack must be <= 3")
        if int(item["sell_price"]) < 0:
            raise ValueError(f"items/item_id {item_id} sell_price must be >= 0")
        if float(item["unit_weight"]) < 0:
            raise ValueError(
                f"items/item_id {item_id} unit_weight cannot be negative"
            )
        if float(item["drop_weight"]) <= 0:
            raise ValueError(
                f"items/item_id {item_id} drop_weight must be > 0"
            )
        if int(item["grid_width"]) != 1 or int(item["grid_height"]) != 1:
            raise ValueError(
                f"items/item_id {item_id} must use a 1x1 grid"
            )

    for active_item_id, item in active_items.items():
        if active_item_id < 10000 or active_item_id >= 20000:
            raise ValueError(
                f"active_item_id {active_item_id} must use 10000..19999"
            )
        if int(item["item_equality"]) not in qualities:
            raise ValueError(
                f"active_item_id {active_item_id} has invalid quality"
            )
        if int(item["max_carry_count"]) <= 0:
            raise ValueError(
                f"active_item_id {active_item_id} max_carry_count must be > 0"
            )
        if int(item["sell_price"]) < 0:
            raise ValueError(
                f"active_item_id {active_item_id} sell_price must be >= 0"
            )
        if float(item["unit_weight"]) < 0:
            raise ValueError(
                f"active_item_id {active_item_id} weight cannot be negative"
            )
        if float(item["drop_weight"]) <= 0:
            raise ValueError(
                f"active_item_id {active_item_id} drop_weight must be > 0"
            )
        if int(item["grid_width"]) != 1 or int(item["grid_height"]) != 1:
            raise ValueError(
                f"active_item_id {active_item_id} must use a 1x1 grid"
            )

    quality_fields = (
        "white_chance",
        "green_chance",
        "blue_chance",
        "purple_chance",
        "gold_chance",
    )
    category_fields = (
        "normal_item_chance",
        "active_item_chance",
        "card_chance",
    )
    for group_id, group in groups.items():
        if int(group["slot_count"]) <= 0:
            raise ValueError(
                f"loot_group_id {group_id} slot_count must be > 0"
            )
        fill_chance = float(group["fill_chance"])
        if fill_chance < 0 or fill_chance > 100:
            raise ValueError(
                f"loot_group_id {group_id} fill_chance must be 0..100"
            )
        chances = [float(group[field]) for field in quality_fields]
        if any(chance < 0 for chance in chances):
            raise ValueError(
                f"loot_group_id {group_id} has a negative quality chance"
            )
        if abs(sum(chances) - 100.0) > 0.001:
            raise ValueError(
                f"loot_group_id {group_id} quality chances must total 100"
            )
        category_chances = [float(group[field]) for field in category_fields]
        if any(chance < 0 for chance in category_chances):
            raise ValueError(
                f"loot_group_id {group_id} has a negative category chance"
            )
        if abs(sum(category_chances) - 100.0) > 0.001:
            raise ValueError(
                f"loot_group_id {group_id} category chances must total 100"
            )
        for quality, chance in enumerate(chances):
            if chance <= 0:
                continue
            has_candidate = any(
                int(item["item_equality"]) == quality
                and float(item["drop_weight"]) > 0
                for item in items.values()
            )
            if not has_candidate:
                raise ValueError(
                    f"loot_group_id {group_id} has chance for quality "
                    f"{quality} but no weighted item candidate"
                )

    actors = rows_to_records(tables["player_npc.csv"])
    for actor in actors:
        item_ids = str(actor["initial_item_ids"]).split(";")
        item_counts = str(actor["initial_item_counts"]).split(";")
        if str(actor["initial_item_ids"]) == "":
            item_ids = []
        if str(actor["initial_item_counts"]) == "":
            item_counts = []
        if len(item_ids) != len(item_counts):
            raise ValueError(
                f"actor_id {actor['actor_id']} initial item arrays differ"
            )
        for item_id, count in zip(item_ids, item_counts, strict=True):
            if int(item_id) not in items:
                raise ValueError(
                    f"actor_id {actor['actor_id']} references item {item_id}"
                )
            if int(count) <= 0:
                raise ValueError("initial item counts must be > 0")


def main() -> None:
    converted: list[tuple[str, int]] = []
    table_rows: dict[str, list[list[Any]]] = {}
    for spec in TABLES:
        rows = read_table(spec)
        table_rows[spec.output] = rows
    validate_card_tables(table_rows)
    validate_loot_tables(table_rows)
    for spec in TABLES:
        rows = table_rows[spec.output]
        write_csv(OUTPUT_DIR / spec.output, rows)
        converted.append((spec.output, max(0, len(rows) - 3)))
    for legacy_output in LEGACY_OUTPUTS:
        (OUTPUT_DIR / legacy_output).unlink(missing_ok=True)
    print("Table conversion complete:")
    for filename, data_rows in converted:
        print(f"  {filename}: {data_rows} data rows")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as error:
        print(f"Table conversion failed: {error}", file=sys.stderr)
        raise SystemExit(1)
